# -*- coding: utf-8 -*-
"""接口级端到端测试：写操作直连 AD + 单条刷新缓存；读操作全部走 SQLite。
严格清理测试数据，避免污染域环境。
"""
import json
import os
import sys
import time
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as app_mod
import config
from modules import db, ad_ldap as L, sync_task

application = app_mod.create_app()
client = application.test_client()

result = []


def rec(name, ok, msg=''):
    result.append((name, ok, msg))
    print(('OK  ' if ok else 'FAIL') + f' {name}' + (f'  -> {msg}' if msg else ''))


def api(method, url, **kw):
    kw.setdefault('data', {})
    r = getattr(client, method)(url, **kw)
    try:
        j = r.get_json()
    except Exception:
        j = None
    return r.status_code, j


# ---------------- 登录
code, j = api('post', '/login', data={'username': 'admin', 'password': config.DEFAULT_ADMIN_PASSWORD})
rec('登录', code == 200 and j and j.get('ok'))

dm0 = db.query_one('SELECT * FROM ad_domain_conf ORDER BY id LIMIT 1')
DID = dm0['id']
base_dn = dm0['base_dn']
print(f'测试域：{dm0["domain_name"]} base={base_dn}')

# ---------------- 只读接口
code, j = api('get', f'/users/api/list?domain_id={DID}&page_size=3')
rec('用户列表接口', code == 200 and j and j.get('ok') and len(j['data']['rows']) <= 3)

code, j = api('get', f'/ous/api/options?domain_id={DID}')
rec('OU 下拉接口', code == 200 and j and len(j['data']) > 1)

code, j = api('get', f'/groups/api/options?domain_id={DID}')
rec('组下拉接口', code == 200 and j and j.get('ok') and len(j['data']) > 0,
    f'{len(j["data"]) if j and j.get("ok") else 0} 个组')

code, j = api('get', f'/users/facets?domain_id={DID}')
rec('表单候选接口(facets)', code == 200 and j and 'departments' in j.get('data', {}))

code, j = api('get', f'/ous/api/tree?domain_id={DID}')
rec('OU 树接口', code == 200 and j and j.get('ok'))

code, j = api('post', f'/domains/{DID}/test')
rec('域连通测试(LDAP)', code == 200 and j and j.get('ok'), (j or {}).get('message', ''))

code, j = api('get', '/domains/api/status')
rec('域状态接口', code == 200)

code, j = api('get', '/domains/api/tasks')
rec('后台任务接口', code == 200)

# ---------------- 写：OU 创建
tag = uuid.uuid4().hex[:6]
ou_name = f'QA-自动化{tag}'
code, j = api('post', '/ous/create', data={'domain_id': DID, 'parent_dn': base_dn,
                                          'name': ou_name, 'description': '自动化测试'})
ok_create_ou = code == 200 and j and j.get('ok')
rec('创建 OU（AD 直连）', ok_create_ou, (j or {}).get('message', ''))
ou_dn = (j or {}).get('data', {}).get('dn', '') if ok_create_ou else ''
if ou_dn:
    row = db.query_one('SELECT * FROM ad_ou_cache WHERE domain_id=? AND dn=?', (DID, ou_dn))
    rec('创建 OU 后缓存单条写入', bool(row))

# ---------------- 写：用户创建
sam = f'qa{tag}'
code, j = api('post', '/users/create', data={
    'domain_id': DID, 'sam_account_name': sam, 'display_name': f'测试用户{tag}',
    'password': 'Abc@12345', 'ou_dn': ou_dn or base_dn, 'department': 'QA',
    'must_change_password': '1', 'enabled': '1', 'mail': f'{sam}@test.local',
})
ok_create_user = code == 200 and j and j.get('ok')
rec('创建用户（AD 直连 + 设密）', ok_create_user, (j or {}).get('message', '')[:120])
user_dn = (j or {}).get('data', {}).get('dn', '') if ok_create_user else ''
uid = None
if user_dn:
    row = db.query_one('SELECT * FROM ad_user_cache WHERE domain_id=? AND dn=?', (DID, user_dn))
    rec('创建用户后缓存单条写入', bool(row), f'enabled={row["enabled"] if row else "-"}')
    uid = row['id'] if row else None
    # AD 侧核实
    conn = L.get_connection(L.get_domain(DID))
    found = conn.search(user_dn, '(objectClass=*)', search_scope='BASE', attributes=['sAMAccountName'])
    rec('AD 侧对象存在', bool(found))

# ---------------- 详情页实时字段
if uid:
    code, j = api('get', f'/users/{uid}')
    rec('用户详情页（实时单条）', code == 200)

# ---------------- 编辑（改名会触发 DN 变更）
if uid:
    new_display = f'测试用户{tag}改'
    code, j = api('post', f'/users/{uid}/update', data={'display_name': new_display, 'title': 'QA工程师'})
    ok = code == 200 and j and j.get('ok')
    new_dn = (j or {}).get('data', {}).get('dn', '')
    rec('编辑用户（重命名→DN 变更）', ok, new_dn)
    if ok:
        cnt = db.scalar('SELECT COUNT(*) FROM ad_user_cache WHERE domain_id=? AND sam_account_name=?', (DID, sam), 0)
        rec('重命名后缓存无重复行', cnt == 1, f'count={cnt}')
        r2 = db.query_one('SELECT * FROM ad_user_cache WHERE domain_id=? AND dn=?', (DID, new_dn))
        rec('重命名后缓存 DN 已更新', bool(r2))
        if r2:
            uid = r2['id']   # 重命名重建缓存行（新自增 id），后续按新 id 操作
        user_dn = new_dn

# ---------------- 禁用 / 启用
if uid:
    code, j = api('post', f'/users/{uid}/toggle', data={'enabled': '0'})
    row = db.query_one('SELECT enabled FROM ad_user_cache WHERE id=?', (uid,))
    rec('禁用账号（AD + 缓存）', code == 200 and j.get('ok') and row['enabled'] == 0)
    code, j = api('post', f'/users/{uid}/toggle', data={'enabled': '1'})
    row = db.query_one('SELECT enabled FROM ad_user_cache WHERE id=?', (uid,))
    rec('启用账号（AD + 缓存）', code == 200 and j.get('ok') and row['enabled'] == 1)

# ---------------- 重置密码（需要 confirm）
if uid:
    code, j = api('post', f'/users/{uid}/reset-password', data={'password': 'Xyz@98765', 'must_change': '1'})
    rec('重置密码缺少二次确认 → 拒绝', code != 200 and j and not j.get('ok'),
        (j or {}).get('message', '')[:60])
    code, j = api('post', f'/users/{uid}/reset-password',
                  data={'password': 'Xyz@98765', 'must_change': '1', 'confirm': '1'})
    rec('重置密码（confirm=1，加密通道）', code == 200 and j and j.get('ok'), (j or {}).get('message', ''))

# ---------------- 移动 OU
if uid and ou_dn:
    code, j = api('post', f'/users/{uid}/move', data={'target_ou': base_dn})
    ok = code == 200 and j and j.get('ok')
    rec('移动用户到域根', ok, (j or {}).get('message', ''))
    if ok:
        row = db.query_one('SELECT id, ou_dn, dn FROM ad_user_cache WHERE domain_id=? AND sam_account_name=?', (DID, sam))
        rec('移动后缓存 ou_dn 正确', row and row['ou_dn'] == base_dn, row['ou_dn'] if row else '-')
        if row:
            uid = row['id']   # 移动同样重建缓存行，按 sam 重新定位
            user_dn = row['dn']

# ---------------- 解锁
if uid:
    code, j = api('post', f'/users/{uid}/unlock')
    rec('解锁账号', code == 200 and j and j.get('ok'))

# ---------------- 用户组：建组 → 加人 → 移人 → 删组
gname = f'QA组{tag}'
code, j = api('post', '/groups/create', data={
    'domain_id': DID, 'name': gname, 'sam_account_name': f'g{tag}',
    'scope': 'Global', 'category': 'Security', 'ou_dn': ou_dn or base_dn,
    'description': '自动化测试组'})
ok_g = code == 200 and j and j.get('ok')
rec('创建用户组', ok_g, (j or {}).get('message', ''))
g_dn = (j or {}).get('data', {}).get('dn', '') if ok_g else ''
gid = None
if g_dn:
    grow = db.query_one('SELECT * FROM ad_group_cache WHERE domain_id=? AND dn=?', (DID, g_dn))
    rec('创建组后缓存写入', bool(grow))
    gid = grow['id'] if grow else None

if gid and user_dn:
    code, j = api('post', f'/groups/{gid}/members', data={'action': 'add', 'member_dns': [user_dn]})
    rec('组添加成员', code == 200 and j and j.get('ok'), (j or {}).get('message', ''))
    code, j = api('post', f'/groups/{gid}/members', data={'action': 'remove', 'member_dns': [user_dn]})
    rec('组移除成员', code == 200 and j and j.get('ok'))

# ---------------- 批量操作（同域）
if uid:
    code, j = api('post', '/users/batch', data={'action': 'disable', 'ids': [uid]})
    rec('批量禁用（单域）', code == 200 and j and j.get('ok'), (j or {}).get('message', ''))
    code, j = api('post', '/users/batch', data={'action': 'enable', 'ids': [uid]})
    rec('批量启用（单域）', code == 200 and j and j.get('ok'))

# ---------------- 导出
r = client.get(f'/users/export?scope=current&domain_id={DID}')
rec('导出用户 Excel', r.status_code == 200 and len(r.data) > 4000, f'{len(r.data)} bytes')
r = client.get(f'/logs/export?domain_id={DID}')
rec('导出日志 Excel', r.status_code == 200 and len(r.data) > 2000, f'{len(r.data)} bytes')
r = client.get(f'/computers/export?domain_id={DID}')
rec('导出计算机 Excel', r.status_code == 200 and len(r.data) > 1000, f'{len(r.data)} bytes')
r = client.get('/users/import/template')
rec('下载导入模板', r.status_code == 200 and len(r.data) > 3000)

# ---------------- 导入校验（用模板回填一行，验证完整链路）
import io
from openpyxl import load_workbook

wb = load_workbook(io.BytesIO(r.data))
ws = wb.active
print('模板表头:', [c.value for c in ws[1]])
imp_sam = f'imp{tag}'
ws.append([imp_sam, f'导入用户{tag}', 'Abc@12345', base_dn, '', '', '', 'QA', '', '导入测试'])
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
r2 = client.post('/users/import/validate',
                 data={'file': (buf, 'test.xlsx'), 'domain_ids': str(DID)},
                 content_type='multipart/form-data')
j2 = r2.get_json()
rec('导入预校验', r2.status_code == 200 and j2 and j2.get('ok'),
    (j2 or {}).get('message', '') if j2 else '')
token = (j2 or {}).get('data', {}).get('token') if j2 and j2.get('ok') else None
if token:
    code, j3 = api('post', '/users/import/run', data={'token': token, 'confirm': '1'})
    okimp = code == 200 and j3 and j3.get('ok')
    rec('导入执行（写 AD + 刷新缓存）', okimp, (j3 or {}).get('message', ''))
    if okimp:
        row = db.query_one('SELECT * FROM ad_user_cache WHERE domain_id=? AND sam_account_name=?', (DID, imp_sam))
        rec('导入后缓存在位', bool(row))
        if row:
            uid_imp = row['id']
            code, j = api('post', f'/users/{uid_imp}/delete', data={'confirm': '1'})
            rec('清理导入用户', code == 200 and j and j.get('ok'))

# ---------------- 清理：删用户 → 删组 → 删 OU
if uid:
    code, j = api('post', f'/users/{uid}/delete', data={'confirm': '1'})
    ok = code == 200 and j and j.get('ok')
    rec('删除用户（AD + 缓存）', ok, (j or {}).get('message', ''))
    if ok:
        cnt = db.scalar('SELECT COUNT(*) FROM ad_user_cache WHERE domain_id=? AND sam_account_name=?', (DID, sam), 0)
        rec('删除后缓存行清除', cnt == 0, f'count={cnt}')

if gid:
    code, j = api('post', f'/groups/{gid}/delete', data={'confirm': '1'})
    ok = code == 200 and j and j.get('ok')
    rec('删除用户组', ok, (j or {}).get('message', ''))
    if ok:
        cnt = db.scalar('SELECT COUNT(*) FROM ad_group_cache WHERE domain_id=? AND dn=?', (DID, g_dn), 0)
        rec('删除组后缓存清除', cnt == 0)

if ou_dn:
    code, j = api('post', '/ous/delete', data={'domain_id': DID, 'dn': ou_dn, 'recursive': '1', 'confirm': '1'})
    ok = code == 200 and j and j.get('ok')
    rec('递归删除 OU', ok, (j or {}).get('message', ''))
    if ok:
        cnt = db.scalar('SELECT COUNT(*) FROM ad_ou_cache WHERE domain_id=? AND dn=?', (DID, ou_dn), 0)
        rec('删除 OU 后缓存清除', cnt == 0)
        # AD 侧确认
        conn = L.get_connection(L.get_domain(DID))
        rec('AD 侧 OU 已不存在', not conn.search(ou_dn, '(objectClass=*)', search_scope='BASE'))

# ---------------- 高危保护
code, j = api('post', f'/domains/{DID}/delete', data={})
rec('删除域缺少 confirm → 拒绝', code != 200 or not (j or {}).get('ok'))
code, j = api('post', '/trust/create', data={'local_id': DID, 'remote_domain_name': 'nonexist.local',
                                            'remote_dc_host': '10.255.255.1', 'remote_username': 'a',
                                            'remote_password': 'b', 'confirm': '1', 'direction': 'bidirectional'})
rec('信任创建（对方不可达）→ 失败且不改缓存', code != 200 and j and not j.get('ok'),
    (j or {}).get('message', '')[:90])
cnt_trust = db.scalar('SELECT COUNT(*) FROM ad_trust_cache', (), 0)
rec('失败后信任缓存未被写入', cnt_trust == 0, f'count={cnt_trust}')

# ---------------- 汇总
print()
bad = [x for x in result if not x[1]]
print(f'===== {len(result) - len(bad)}/{len(result)} 通过 =====')
for n, ok, m in bad:
    print(f'  FAIL {n}: {m}')
